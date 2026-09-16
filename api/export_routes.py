"""XLSX export of a single hunt's leads, for handing results to business users.

Scope: one hunt at a time. The global lead repository is deliberately not
exportable — it accumulates across hunts without bound, so a bulk export is a
memory/latency risk; page it through ``GET /api/v1/leads`` when needed.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any, Callable

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from api.hunt_store import load_hunt
from api.security import require_api_access
from persistence import lead_repo
from persistence.db import run_db

router = APIRouter(prefix="/api/v1", tags=["export"], dependencies=[Depends(require_api_access)])

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_COLUMN_WIDTH = 60

Column = tuple[str, Callable[[dict[str, Any]], Any]]


def _joined(values: Any, separator: str = "; ") -> str:
    if not isinstance(values, (list, tuple)):
        return str(values or "")
    return separator.join(str(item).strip() for item in values if str(item).strip())


def _first_decision_maker(lead: dict[str, Any]) -> dict[str, Any]:
    makers = lead.get("decision_makers") or []
    if isinstance(makers, list):
        for maker in makers:
            if isinstance(maker, dict) and (maker.get("name") or maker.get("email")):
                return maker
    return {}


def _contact_name(lead: dict[str, Any]) -> str:
    contact = str(lead.get("contact_person", "") or "").strip()
    if contact:
        return contact
    return str(_first_decision_maker(lead).get("name", "") or "").strip()


def _all_decision_makers(lead: dict[str, Any]) -> str:
    makers = lead.get("decision_makers") or []
    parts: list[str] = []
    if isinstance(makers, list):
        for maker in makers:
            if not isinstance(maker, dict):
                continue
            name = str(maker.get("name", "") or "").strip()
            title = str(maker.get("title", "") or "").strip()
            email = str(maker.get("email", "") or "").strip()
            label = " ".join(part for part in (name, f"({title})" if title else "") if part)
            if email:
                label = f"{label} <{email}>" if label else email
            if label:
                parts.append(label)
    return " | ".join(parts)


def _social_links(lead: dict[str, Any]) -> str:
    social = lead.get("social_media") or {}
    if not isinstance(social, dict):
        return ""
    return " | ".join(f"{key}: {value}" for key, value in social.items() if value)


def _score(value: Any) -> Any:
    try:
        return round(float(value), 3)
    except (TypeError, ValueError):
        return ""


BRIEF_COLUMNS: list[Column] = [
    ("公司名称", lambda l: l.get("company_name", "")),
    ("官网", lambda l: l.get("website", "")),
    ("国家/地区", lambda l: l.get("country_code", "")),
    ("行业", lambda l: l.get("industry", "")),
    ("联系人", _contact_name),
    ("邮箱", lambda l: _joined(l.get("emails"))),
    ("电话", lambda l: _joined(l.get("phone_numbers"))),
    ("优先级", lambda l: l.get("priority_tier", "")),
    ("匹配度", lambda l: _score(l.get("match_score"))),
]

FULL_COLUMNS: list[Column] = BRIEF_COLUMNS + [
    ("联系人职务", lambda l: _first_decision_maker(l).get("title", "")),
    ("全部决策人", _all_decision_makers),
    ("地址", lambda l: l.get("address", "")),
    ("社交媒体", _social_links),
    ("客户类型", lambda l: l.get("customer_role", "")),
    ("可触达度", lambda l: _score(l.get("contactability_score"))),
    ("契合度", lambda l: _score(l.get("fit_score"))),
    ("证据强度", lambda l: l.get("evidence_strength", "")),
    ("来源关键词", lambda l: l.get("source_keyword", "")),
    ("首次发现", lambda l: l.get("first_seen_at", "")),
    ("最近更新", lambda l: l.get("last_seen_at", "")),
    ("复用线索", lambda l: "是" if l.get("reused") else "否"),
]

VIEWS: dict[str, list[Column]] = {"brief": BRIEF_COLUMNS, "full": FULL_COLUMNS}


def build_workbook(leads: list[dict[str, Any]], columns: list[Column], *, sheet_title: str) -> bytes:
    """Render leads into a single-sheet xlsx and return the bytes."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title[:31] or "Leads"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    for index, (header, _) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for row_index, lead in enumerate(leads, start=2):
        for col_index, (_, extractor) in enumerate(columns, start=1):
            try:
                value = extractor(lead)
            except Exception:  # noqa: BLE001 - a bad row must not fail the export
                value = ""
            if isinstance(value, (dict, list)):
                value = str(value)
            elif value is None:
                value = ""
            sheet.cell(row=row_index, column=col_index, value=value)

    for col_index, (header, _) in enumerate(columns, start=1):
        longest = len(header)
        for row_index in range(2, len(leads) + 2):
            cell = sheet.cell(row=row_index, column=col_index)
            if cell.value is not None:
                longest = max(longest, len(str(cell.value)))
        sheet.column_dimensions[get_column_letter(col_index)].width = min(longest + 2, MAX_COLUMN_WIDTH)

    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@router.get("/hunts/{hunt_id}/export")
async def export_hunt_leads(
    hunt_id: str,
    format: str = Query(default="xlsx", description="Export format; only xlsx is supported"),
    view: str = Query(default="brief", description="brief | full"),
):
    """Download one hunt's leads as an xlsx workbook (brief or full columns)."""
    if format.lower() != "xlsx":
        raise HTTPException(status_code=400, detail="Only format=xlsx is supported")
    columns = VIEWS.get(view.lower())
    if columns is None:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported view '{view}'; expected one of: {', '.join(sorted(VIEWS))}",
        )

    if not await run_db(lambda: load_hunt(hunt_id)):
        raise HTTPException(status_code=404, detail="Hunt not found")

    leads = await run_db(lead_repo.list_hunt_leads, hunt_id)
    content = build_workbook(leads, columns, sheet_title=view.lower())
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d")
    filename = f"hunt-{hunt_id[:8]}-{view.lower()}-{stamp}.xlsx"
    return Response(
        content=content,
        media_type=XLSX_MEDIA_TYPE,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "X-Lead-Count": str(len(leads)),
            "X-Export-View": view.lower(),
        },
    )
