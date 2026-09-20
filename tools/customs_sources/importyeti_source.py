"""ImportYeti ingestion — normalize bill-of-lading exports into TradeRecords.

ImportYeti (importyeti.com) republishes US Customs AMS bills of lading (FOIA).
Its CSV/XLSX exports carry shipment rows where the consignee is the US
importer — i.e. our prospect — and the shipper is the supplier.

Column headers vary across export formats and versions, so headers are matched
against alias lists; an explicit ``column_mapping`` (field -> header) always
wins and is the fallback when detection fails.
"""

from __future__ import annotations

import csv
import hashlib
import logging
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

logger = logging.getLogger(__name__)

SOURCE_NAME = "importyeti_csv"

# Candidate headers per canonical field, most specific first. Matching is done
# on normalized (lowercase, single-spaced) headers, exact match before contains.
_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "consignee_name": (
        "consignee name", "consignee", "importer name", "importer",
        "buyer name", "buyer", "company name", "company", "receiver",
    ),
    "supplier_name": (
        "shipper name", "shipper", "supplier name", "supplier",
        "exporter name", "exporter", "seller", "vendor",
    ),
    "country": (
        "supplier country", "shipper country", "origin country",
        "country of origin", "pod country", "destination country", "country",
    ),
    "hs_code": (
        "hs code", "hscode", "hs codes", "hts code", "hts", "hts number",
        "hsn", "tariff code", "commodity code",
    ),
    "arrival_date": (
        "arrival date", "arrival", "shipment date", "bill of lading date",
        "bol date", "dated", "eta", "date",
    ),
    "quantity": (
        "quantity", "qty", "units", "pieces", "cartons", "packages", "containers",
    ),
    "weight": ("gross weight", "net weight", "weight kg", "weight"),
    "product_description": (
        "product description", "goods description", "cargo description",
        "commodity description", "goods shipped", "commodity", "merchandise",
        "description", "goods",
    ),
}

# US-format dates first: ImportYeti data is US imports.
_DATE_FORMATS = (
    "%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d", "%m-%d-%Y",
    "%b %d, %Y", "%B %d, %Y", "%d %b %Y", "%d %B %Y", "%d-%b-%Y",
    "%m/%d/%y", "%d-%b-%y",
)

_COUNTRY_ISO2 = {
    "united states": "us", "usa": "us", "u.s.a.": "us", "china": "cn",
    "mainland china": "cn", "hong kong": "hk", "taiwan": "tw",
    "vietnam": "vn", "viet nam": "vn", "india": "in", "south korea": "kr",
    "korea": "kr", "japan": "jp", "thailand": "th", "indonesia": "id",
    "malaysia": "my", "singapore": "sg", "philippines": "ph",
    "turkey": "tr", "germany": "de", "france": "fr", "italy": "it",
    "spain": "es", "poland": "pl", "netherlands": "nl", "belgium": "be",
    "united kingdom": "gb", "uk": "gb", "mexico": "mx", "brazil": "br",
    "canada": "ca", "bangladesh": "bd", "pakistan": "pk", "cambodia": "kh",
    "sri lanka": "lk", "uae": "ae", "israel": "il", "guatemala": "gt",
    "honduras": "hn", "el salvador": "sv", "peru": "pe", "colombia": "co",
}

_ISO2_RE = re.compile(r"^[a-z]{2}$")


def normalize_header(header: str) -> str:
    return " ".join(str(header or "").strip().lower().split())


def detect_column_mapping(headers: list[str]) -> dict[str, str]:
    """Map canonical fields to actual headers via exact-then-contains alias match."""
    normalized = {normalize_header(h): h for h in headers if normalize_header(h)}
    mapping: dict[str, str] = {}
    for field_name, aliases in _COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                mapping[field_name] = normalized[alias]
                break
        else:
            for alias in aliases:
                for norm, original in normalized.items():
                    if alias in norm:
                        mapping[field_name] = original
                        break
                if field_name in mapping:
                    break
    return mapping


def normalize_date(value: object) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    iso = re.match(r"^(\d{4})-(\d{2})-(\d{2})", raw)
    if iso:
        return iso.group(0)
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(raw, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return ""


def normalize_country_code(value: object) -> str:
    raw = normalize_header(str(value or ""))
    if not raw:
        return ""
    if _ISO2_RE.match(raw):
        return raw
    return _COUNTRY_ISO2.get(raw, "")


def normalize_hs_code(value: object) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    if len(digits) in (6, 8, 10):
        return digits
    return ""


def _clean(value: object) -> str:
    return " ".join(str(value or "").split())


@dataclass
class TradeRecord:
    """One normalized shipment row, source-agnostic."""

    company_name: str          # canonical prospect name (= consignee for imports)
    consignee_name: str
    supplier_name: str
    country: str               # supplier country display name
    country_code: str          # supplier country ISO2
    hs_code: str
    product_description: str
    arrival_date: str          # ISO YYYY-MM-DD or ""
    quantity: str
    weight: str
    source: str = SOURCE_NAME
    source_ref: str = ""
    domain: str = ""
    raw: dict = field(default_factory=dict)
    record_hash: str = ""

    def finalize(self) -> "TradeRecord":
        if not self.record_hash:
            # source_ref (file name) is deliberately excluded so the same
            # shipment row arriving via overlapping exports still dedupes.
            basis = "|".join(
                (
                    self.company_name.lower(),
                    self.supplier_name.lower(),
                    self.arrival_date,
                    self.hs_code,
                    self.quantity,
                    self.weight,
                    self.source,
                )
            )
            self.record_hash = hashlib.sha1(basis.encode("utf-8")).hexdigest()
        return self


def record_from_row(
    row: dict,
    *,
    column_mapping: dict[str, str] | None = None,
    source_ref: str = "",
) -> TradeRecord | None:
    """Convert one raw row dict into a TradeRecord; None when unusable."""
    mapping = column_mapping or detect_column_mapping(list(row.keys()))

    def pick(field_name: str) -> str:
        header = mapping.get(field_name)
        return _clean(row.get(header)) if header else ""

    consignee = pick("consignee_name")
    if not consignee:
        return None  # the consignee IS the prospect; without it the row is noise

    supplier = pick("supplier_name")
    country = pick("country")
    return TradeRecord(
        company_name=consignee,
        consignee_name=consignee,
        supplier_name=supplier,
        country=country,
        country_code=normalize_country_code(country),
        hs_code=normalize_hs_code(pick("hs_code")),
        product_description=pick("product_description"),
        arrival_date=normalize_date(pick("arrival_date")),
        quantity=pick("quantity"),
        weight=pick("weight"),
        source_ref=source_ref,
        raw=dict(row),
    ).finalize()


def load_rows(path: str | Path) -> list[dict]:
    """Load a CSV or XLSX export into a list of raw row dicts."""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _load_xlsx(p)
    if suffix == ".xls":
        raise ValueError(f"legacy .xls is not supported, re-export as CSV/XLSX: {p.name}")
    return _load_csv(p)


def _load_csv(p: Path) -> list[dict]:
    text: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = p.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError(f"cannot decode file: {p.name}")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(text.splitlines(), dialect=dialect)
    return [dict(row) for row in reader if any(_clean(v) for v in row.values() if v)]


def _load_xlsx(p: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(p, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers: list[str] = []
        out: list[dict] = []
        for values in rows:
            cells = ["" if v is None else str(v) for v in values]
            if not any(c.strip() for c in cells):
                continue
            if not headers:
                headers = cells
                continue
            row = dict(zip(headers, cells))
            if any(_clean(v) for v in row.values()):
                out.append(row)
        return out
    finally:
        wb.close()


def parse_file(
    path: str | Path,
    *,
    column_mapping: dict[str, str] | None = None,
    source_ref: str = "",
) -> list[TradeRecord]:
    """Parse an ImportYeti export file into normalized TradeRecords."""
    rows = load_rows(path)
    mapping = column_mapping or (
        detect_column_mapping(list(rows[0].keys())) if rows else {}
    )
    records: list[TradeRecord] = []
    for row in rows:
        try:
            item = record_from_row(row, column_mapping=mapping, source_ref=source_ref or Path(path).name)
        except Exception:
            logger.debug("[ImportYetiSource] dropped malformed row in %s", Path(path).name, exc_info=True)
            continue
        if item:
            records.append(item)
    logger.info("[ImportYetiSource] parsed %d/%d rows from %s", len(records), len(rows), Path(path).name)
    return records
