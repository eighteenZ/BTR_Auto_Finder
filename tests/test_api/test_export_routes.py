"""Tests for the per-hunt xlsx export endpoint."""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.hunter_app import create_hunter_app

LEADS = [
    {
        "company_name": "Acme Importers",
        "website": "https://acme.com",
        "country_code": "US",
        "industry": "Wholesale",
        "contact_person": "",
        "emails": ["buyer@acme.com", "info@acme.com"],
        "phone_numbers": ["+1 555 0100"],
        "priority_tier": "high",
        "match_score": 0.8123,
        "fit_score": 0.7,
        "contactability_score": 0.66,
        "address": "1 Main St, Dallas, TX",
        "social_media": {"linkedin": "https://linkedin.com/company/acme"},
        "decision_makers": [
            {"name": "Jane Doe", "title": "Purchasing Manager", "email": "jane@acme.com"},
            {"name": "Bob Ray", "title": "Owner", "email": ""},
        ],
        "customer_role": "importer",
        "source_keyword": "wholesale importer",
        "first_seen_at": "2026-09-15T10:00:00Z",
        "last_seen_at": "2026-09-16T09:00:00Z",
        "evidence_strength": "medium",
        "reused": False,
    },
    {
        "company_name": "Beta Trading",
        "website": "https://beta.example",
        "country_code": "US",
        "industry": "Distribution",
        "contact_person": "Ann Lee",
        "emails": [],
        "phone_numbers": [],
        "priority_tier": "low",
        "match_score": 0.31,
        "decision_makers": [],
        "reused": True,
    },
]


def _client(monkeypatch, *, hunt_exists: bool = True, leads: list | None = None) -> TestClient:
    monkeypatch.setattr("api.export_routes.load_hunt", lambda hunt_id: ({"status": "completed"} if hunt_exists else None))
    monkeypatch.setattr("api.export_routes.lead_repo.list_hunt_leads", lambda hunt_id: list(leads if leads is not None else LEADS))
    return TestClient(create_hunter_app())


def _sheet(response):
    workbook = load_workbook(io.BytesIO(response.content))
    return workbook.active


class TestExportHuntLeads:
    def test_brief_export_returns_xlsx_attachment(self, monkeypatch):
        client = _client(monkeypatch)
        res = client.get("/api/v1/hunts/hunt-12345678/export?view=brief")

        assert res.status_code == 200
        assert res.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert 'filename="hunt-hunt-123-brief-' in res.headers["content-disposition"]
        assert res.headers["x-lead-count"] == "2"

        sheet = _sheet(res)
        headers = [c.value for c in sheet[1]]
        assert headers == [
            "公司名称", "官网", "国家/地区", "行业", "联系人",
            "邮箱", "电话", "优先级", "匹配度",
        ]
        assert sheet["A2"].value == "Acme Importers"
        assert sheet["D2"].value == "Wholesale"
        assert sheet["E2"].value == "Jane Doe"          # 无 contact_person 时取首位决策人
        assert sheet["F2"].value == "buyer@acme.com; info@acme.com"
        assert sheet["G2"].value == "+1 555 0100"
        assert sheet["I2"].value == 0.812                # 分数四舍五入到 3 位
        assert sheet["E3"].value == "Ann Lee"            # 有 contact_person 时优先用它
        assert sheet.freeze_panes == "A2"

    def test_full_view_adds_business_and_evidence_columns(self, monkeypatch):
        client = _client(monkeypatch)
        res = client.get("/api/v1/hunts/hunt-12345678/export?view=full")

        assert res.status_code == 200
        sheet = _sheet(res)
        headers = [c.value for c in sheet[1]]
        assert headers[:9] == [
            "公司名称", "官网", "国家/地区", "行业", "联系人",
            "邮箱", "电话", "优先级", "匹配度",
        ]
        for expected in ("全部决策人", "社交媒体", "来源关键词", "复用线索", "可触达度"):
            assert expected in headers

        row = {header: sheet.cell(row=2, column=i + 1).value for i, header in enumerate(headers)}
        assert row["联系人职务"] == "Purchasing Manager"
        assert "Jane Doe (Purchasing Manager) <jane@acme.com>" in row["全部决策人"]
        assert "Bob Ray (Owner)" in row["全部决策人"]
        assert row["社交媒体"] == "linkedin: https://linkedin.com/company/acme"
        assert row["来源关键词"] == "wholesale importer"
        assert row["复用线索"] == "否"
        assert sheet.cell(row=3, column=headers.index("复用线索") + 1).value == "是"

    def test_empty_hunt_still_returns_headers(self, monkeypatch):
        client = _client(monkeypatch, leads=[])
        res = client.get("/api/v1/hunts/hunt-empty/export")

        assert res.status_code == 200
        assert res.headers["x-lead-count"] == "0"
        assert [c.value for c in _sheet(res)[1]][0] == "公司名称"

    def test_default_view_is_full(self, monkeypatch):
        client = _client(monkeypatch)
        res = client.get("/api/v1/hunts/hunt-12345678/export")

        assert res.status_code == 200
        assert res.headers["x-export-view"] == "full"
        sheet = _sheet(res)
        headers = [c.value for c in sheet[1]]
        assert headers[:10] == [
            "公司名称", "官网", "国家/地区", "行业", "联系人", "邮箱", "电话",
            "优先级", "匹配度", "域名",
        ]
        # “全量”=全部字段：包含海关与证据等完整列
        for expected in ("海关数据评分", "海关数据", "线索ID", "线索键"):
            assert expected in headers

    def test_oversized_cell_is_truncated_not_fatal(self, monkeypatch):
        huge = dict(LEADS[0])
        huge["customs_data"] = "X" * 60000
        client = _client(monkeypatch, leads=[huge])

        res = client.get("/api/v1/hunts/hunt-1/export?view=full")
        assert res.status_code == 200
        sheet = _sheet(res)
        customs_col = [c.value for c in sheet[1]].index("海关数据")
        cell = sheet.cell(row=2, column=customs_col + 1).value
        assert len(cell) <= 32767 and cell.endswith("…[truncated]")

    def test_unknown_hunt_returns_404(self, monkeypatch):
        client = _client(monkeypatch, hunt_exists=False)
        assert client.get("/api/v1/hunts/missing/export").status_code == 404

    @pytest.mark.parametrize("query", ["format=csv", "view=summary", "view="])
    def test_invalid_parameters_return_400(self, monkeypatch, query):
        client = _client(monkeypatch)
        assert client.get(f"/api/v1/hunts/hunt-1/export?{query}").status_code == 400

    def test_bad_row_does_not_fail_the_export(self, monkeypatch):
        broken = dict(LEADS[0])
        broken["social_media"] = "not-a-dict"
        client = _client(monkeypatch, leads=[broken])
        res = client.get("/api/v1/hunts/hunt-1/export?view=full")
        assert res.status_code == 200
